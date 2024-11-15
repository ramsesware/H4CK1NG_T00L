from PyPDF2 import PdfReader, PdfWriter
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from tkinter import messagebox
import os
from PIL import Image
import piexif
import tkinter as tk
import zipfile
import shutil
from mutagen import File as MutagenFile
from hachoir.parser import createParser
from hachoir.metadata import extractMetadata
from hachoir.stream import FileOutputStream
from hachoir.editor import createEditor
import xml.etree.ElementTree as ET

def analyze_metadata(filepath):
    try:
        if filepath.endswith('.pdf'):
            with open(filepath, 'rb') as file:
                pdf = PdfReader(file)
                info = pdf.metadata
                return info
        elif filepath.endswith('.docx'):
            doc = Document(filepath)
            props = doc.core_properties
            return {
                "Identifier": props.identifier or "N/A",
                "Title": props.title or "N/A",
                "Subject": props.subject or "N/A",
                "Author": props.author or "N/A",
                "Last author": props.last_modified_by or "N/A",
                "Date creation at": props.created or "N/A",
                "Modified at": props.modified or "N/A",
                "Category": props.category or "N/A",
                "Language": props.language or "N/A",
                "Content status": props.content_status or "N/A",
                "Keywords": props.keywords or "N/A",
                "Revision": props.revision or "N/A",
                "Last printed": props.last_printed or "N/A",
                "Comments": props.comments or "N/A",
                "Version": props.version or "N/A"
            }
        elif filepath.endswith('.xlsx'):
            workbook = load_workbook(filepath, read_only=True)
            props = workbook.properties
            if props.creator == "openpyxl":
                props.creator = None
                props.created = None
                props.modified = None
            return {
                "Title": props.title or "N/A",
                "Author": props.creator or "N/A",
                "Last author": props.lastModifiedBy or "N/A",
                "Date creation at": props.created or "N/A",
                "Last modified at": props.modified or "N/A",
                "Category": props.category or "N/A",
                "Comments": props.description or "N/A"
            }
        elif filepath.endswith('.pptx'):
            presentation = Presentation(filepath)
            props = presentation.core_properties
            return {
                "Title": props.title or "N/A",
                "Author": props.creator or "N/A",
                "Last author": props.lastModifiedBy or "N/A",
                "Date creation at": props.created or "N/A",
                "Last modified at": props.modified or "N/A",
                "Category": props.category or "N/A",
                "Comments": props.description or "N/A"
            }
        elif filepath.endswith(('.jpg', '.jpeg', '.png')):
            image = Image.open(filepath)
            if "exif" in image.info:
                exif_data = piexif.load(image.info["exif"])
                metadata = {}
                for ifd in exif_data:
                    for tag, value in exif_data[ifd].items():
                        tag_name = piexif.TAGS[ifd].get(tag, tag)
                        metadata[tag_name] = value
                return metadata
        elif filepath.endswith(('.mp3', '.flac', '.wav', '.ogg')):
            audio = MutagenFile(filepath)
            return audio.tags if audio else "No tags found"
        elif filepath.endswith(('.mp4', '.mkv', '.avi', '.mov')):
            parser = createParser(filepath)
            if not parser:
                return "Unable to parse video file"
            metadata = extractMetadata(parser)
            return metadata.exportDictionary() if metadata else "No metadata found"
    except Exception as e:
        messagebox.showerror("Error", f"Error on file analyzing: {e}")


def remove_metadata_audio(filepath):
    audio = MutagenFile(filepath, easy=True)
    if not audio:
        return f"No metadata found in {filepath}."
    
    audio.delete()
    audio.save()


def remove_metadata_video(filepath):

    parser = createParser(filepath)
    if not parser:
        return f"Unable to parse video file {filepath}."
    
    editor = createEditor(parser)
    if not editor:
        return f"Unable to create editor for {filepath}."
    
    for field in list(editor.iterFields()):
        editor.removeField(field)
    
    output_filepath = filepath.replace(".mp4", "_cleaned.mp4")  # Example for .mp4
    with open(output_filepath, "wb") as output_file:
        stream = FileOutputStream(output_file)
        editor.writeInto(stream)



def remove_metadata_pdf(filepath):
    reader = PdfReader(filepath)
    writer = PdfWriter()
    for page in range(len(reader.pages)):
        writer.add_page(reader.pages[page])

    writer.add_metadata({})

    with open(filepath, "wb") as f:
        writer.write(f)

def remove_metadata_office(filepath):
    temp_dir = "temp_file"
    file_extension = filepath.split('.')[-1]
    
    with zipfile.ZipFile(filepath, 'r') as zip_ref:
        zip_ref.extractall(temp_dir)
    
    metadata_files = {
        'docx': ['docProps/core.xml', 'docProps/app.xml'],
        'xlsx': ['docProps/core.xml', 'docProps/app.xml'],
        'pptx': ['docProps/core.xml', 'docProps/app.xml']
    }
    
    if file_extension not in metadata_files:
        raise ValueError("File format not supported")
    
    for meta_file in metadata_files[file_extension]:
        meta_path = os.path.join(temp_dir, meta_file)
        if os.path.exists(meta_path):
            tree = ET.parse(meta_path)
            root = tree.getroot()
            for elem in root.iter():
                elem.clear()
            tree.write(meta_path)

    with zipfile.ZipFile(filepath, 'w') as zip_ref:
        for folder_name, subfolders, filenames in os.walk(temp_dir):
            for filename in filenames:
                file_path = os.path.join(folder_name, filename)
                zip_ref.write(file_path, os.path.relpath(file_path, temp_dir))

    shutil.rmtree(temp_dir)

def remove_metadata_image(filepath):
    imagen = Image.open(filepath)
    if "exif" in imagen.info:
        imagen.info.pop("exif")
    imagen.save(filepath)

def remove_metadata_file(filepath):
    try:
        extension = os.path.splitext(filepath)[1]
        if extension in ['.jpg', '.jpeg', '.png']:
            remove_metadata_image(filepath)
            return f"File: metadata from {os.path.basename(filepath)} has been removed correctly."
        elif extension == '.pdf':
            remove_metadata_pdf(filepath)
            return f"File: metadata from {os.path.basename(filepath)} has been removed correctly."
        elif extension in ['.docx', '.xlsx', '.pptx']:
            remove_metadata_office(filepath)
            return f"File: metadata from {os.path.basename(filepath)} has been removed correctly."
        elif extension in ['.mp3', '.flac', '.wav', '.ogg']:
            remove_metadata_audio(filepath)
            return f"File: metadata from {os.path.basename(filepath)} has been removed correctly."
        elif extension in ['.mp4', '.mkv', '.avi', '.mov']:
            remove_metadata_video(filepath)
            return f"File: metadata from {os.path.basename(filepath)} has been removed correctly."
        
        else:
            return f"Program doesn't support metadata removing for extension: {extension} of {os.path.basename(filepath)}"
            
    except Exception as e:
        messagebox.showerror("Error", f"Error on removing metadata from a file: {e}")

def remove_metadata_directory(file_list):
    try:
        info_list = []
        for file in file_list:
            info_list.append(remove_metadata_file(file))
        return info_list
    except Exception as e:
        messagebox.showerror("Error", f"Error on removing metadata from a directory: {e}")

def analyze_metadata_directory(file_list):
    try:
        info_list = []
        for file_read in file_list:
            info = analyze_metadata(file_read)
            info_list.append({
                "filename": file_read,
                "metadata": info
            }) 
        return info_list
    except Exception as e:
        messagebox.showerror("Error", f"Error on file analyzing: {e}")

def clear_results_metadata(result_text):
    result_text.delete("1.0", tk.END)
    messagebox.showinfo("Results cleared", "Text area have been cleared.")
